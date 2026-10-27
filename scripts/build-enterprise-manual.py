"""Build the Chinese operation manual from screenshots of the deployed sandbox.

Use a dedicated document runtime with python-docx, Pillow, Noto Sans CJK SC,
LibreOffice, Poppler and pdf2image. Render and inspect every page after generation.
"""
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--screenshots', type=Path, required=True)
parser.add_argument('--output', type=Path, required=True)
parser.add_argument('--url', default='http://192.168.1.11:18860')
parser.add_argument('--version', required=True)
args = parser.parse_args()
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
section.left_margin, section.right_margin = Cm(2), Cm(2)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Number', 'List Bullet', 'Caption']:
    s = doc.styles[name]
    s.font.name = 'Noto Sans CJK SC'
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    s.font.color.rgb = RGBColor.from_string('172C48' if name != 'Title' else '000000')
    s.font.size = Pt(10.5 if name == 'Normal' else 11)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.line_spacing = 1.25
for name, size in [('Title', 27), ('Heading 1', 19), ('Heading 2', 13)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
for name in ['Heading 1', 'Heading 2']:
    doc.styles[name].paragraph_format.keep_with_next = True
# Remove Word template borders, including the inherited Title blue rule.
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('企业AI 智能运维平台  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
footer.style = doc.styles['Caption']
doc.core_properties.title = '企业AI 智能运维平台操作手册'
doc.core_properties.subject = '浏览器演示沙盘操作说明'
doc.core_properties.author = 'Enterprise AI'
doc.core_properties.keywords = '演示 运维 操作手册'

def paragraph(text):
    doc.add_paragraph(text)

def steps(items):
    numbering = doc.part.numbering_part.element
    base_id = doc.styles['List Number'].element.pPr.numPr.numId.val
    base = numbering.xpath(f'./w:num[@w:numId="{base_id}"]/w:abstractNumId')[0].val
    num = numbering.add_num(base)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    for item in items:
        p = doc.add_paragraph(item, style='List Number')
        p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num.numId

def screenshot(name, caption):
    path = args.screenshots / (name + '.png')
    if not path.exists():
        raise FileNotFoundError(path)
    doc.add_picture(str(path), width=Cm(17))
    doc.add_paragraph(caption, style='Caption')

def page(title, intro, items, image=None, caption=None, result=None):
    doc.add_page_break()
    doc.add_heading(title, 1)
    paragraph(intro)
    steps(items)
    if image:
        screenshot(image, caption or title)
    if result:
        doc.add_heading('预期结果', 2)
        paragraph(result)

# Cover and static contents are deterministic and readable without field updates.
doc.add_paragraph('企业AI 智能运维平台\n操作手册', style='Title')
doc.add_paragraph('独立定制演示版', style='Subtitle')
paragraph('本手册面向演示操作员和培训人员，说明如何在企业运维工作台中查看资源、处理告警、体验 AI 分析、模拟执行流程及维护知识。按照手册步骤，可以完成从发现风险到记录处置过程的演示闭环。')
paragraph('本系统使用虚构企业数据。AI 回答为预设内容，流程只生成本地模拟记录，不连接真实设备或外部模型，不用于生产操作。')
paragraph(f'访问地址  {args.url}\n构建版本  {args.version}\n进入方式  点击进入演示 无需账号密码')
doc.add_heading('阅读导航', 1)
for line in ['1 访问平台和进入演示', '2 认识工作台与业务导航', '3 查询资源和查看详情', '4 检查模拟监测指标', '5 确认和关闭告警', '6 使用 AI 运维助手', '7 模拟执行自动化流程', '8 查看执行记录', '9 阅读和编辑运维知识', '10 演示高 CPU 处置闭环', '11 演示接口延迟分析', '12 演示每日健康巡检', '13 个性化设置和恢复示例', '14 常见问题与平台说明']:
    paragraph(line)
page('1 访问平台和进入演示', '使用现代桌面浏览器访问测试地址。初次访问默认使用中文浅色界面，页面明确标注浏览器演示沙盘。', ['打开访问地址，确认产品名称和 X 标识。', '阅读数据和模拟操作说明，点击进入演示。', '确认出现运维总览。退出时打开右上角账户菜单，选择退出演示。'], 'login', '图 1 演示入口', '无需注册或密码；退出后本浏览器保存的数据不会删除。共享电脑演示结束后可清理站点数据。')
page('2 认识工作台与业务导航', '顶部按业务组织为运营驾驶舱、风险处置、资产与观测、智能作业、运维知识五个区域。左侧显示当前区域的二级页面。', ['进入运营驾驶舱，查看 24 项资源、88% 健康率、9 项待跟进风险和 6 条初始执行记录。', '在优先处置列表选择告警，或使用下方资源全景、智能巡检和知识沉淀快捷入口。', '窄屏下点击面包屑左侧导航按钮，选择目标页面；按 Escape 可关闭抽屉。'], 'overview', '图 2 运维总览和业务导航', '指标是示例快照。处理告警或执行流程后，对应统计会更新，不代表实时采集结果。')
page('3 查询资源和查看详情', '资源台账按 commerce、finance、platform 三个业务域组织应用主机、数据库和网络设备。', ['进入资产与观测中的资源台账。', '输入 order-api-01 或资源 IP，可再按业务域筛选。无结果时点击清除筛选。', '点击查看详情，核对资源编号、CPU、内存与关联告警。', '点击查看指标监测，跳转到该资源的趋势。'], 'assets', '图 3 资源台账', '示例资源共 24 项。筛选只改变显示结果，不删除数据；窄屏表格可在表格内横向滚动查看末列。')
page('4 检查模拟监测指标', '指标监测提供资源与时间范围切换，显示 CPU 和内存使用率趋势。', ['进入资产与观测中的指标监测。', '在监测资源中选择 order-api-01，可输入名称搜索。', '分别选择近 6 小时、近 12 小时、近 24 小时，观察曲线范围。', '将鼠标移到曲线上查看数值提示。'], 'metrics', '图 4 指标监测', '时间范围只调整模拟趋势，不会向采集端或外部系统请求数据；指标卡片展示固定资源快照。')
page('5 确认和关闭告警', '告警中心包含待处置、处理中和已关闭三种状态，支持名称、编号和状态筛选。', ['进入风险处置，搜索 INC-1201，点击查看详情。', '点击确认告警，将状态从待处置变为处理中。', '点击关闭告警，先选择取消，确认状态未变化。', '再次点击关闭告警并确认，状态变为已关闭。已关闭告警不能再次确认。'], 'incident', '图 5 告警详情与处置入口', '确认、关闭和取消均有不同结果；刷新浏览器后，已保存的状态仍保留。恢复示例可重新演示。')
page('6 使用 AI 运维助手', 'AI 运维助手提供高 CPU 应急排查、接口延迟诊断、每日健康巡检三个预设场景。回答明确标注为模拟分析结果。', ['进入智能作业中的 AI 运维助手。', '点击高 CPU 应急排查，查看 order-api-01、94% CPU 快照和 INC-1201 的关联说明。', '点击查看告警证据或查看参考知识，核对相关页面。', '点击前往关联流程，进入高 CPU 流程。也可输入 CPU、延迟或巡检问题体验关键词路由。'], 'assistant', '图 6 AI 模拟分析', '未覆盖的问题会提示支持范围。本版不调用真实模型，回答不应当作生产诊断结论。')
page('7 模拟执行自动化流程', '三个流程分别用于健康巡检、高 CPU 排查和接口延迟诊断，每个流程包含四个可见步骤。', ['进入智能作业中的自动化流程。', '在目标流程点击模拟执行，阅读确认说明。', '选择取消时不产生记录；再次执行并确认，查看模拟完成提示。', '点击查看记录，核对新增执行。'], 'automation', '图 7 自动化流程', '模拟执行不会访问主机或执行命令，也不会自动关闭告警。按钮在确认期间禁用，避免重复提交。')
page('8 查看执行记录', '执行记录用于演示过程追溯，包含执行编号、流程、时间、步骤数和模拟结果。', ['进入智能作业中的执行记录。', '核对初始六条示例记录，或在模拟流程后查看顶部新增记录。', '刷新浏览器，再确认新增记录仍存在。', '需要重新演示时使用恢复示例数据，记录将回到初始状态。'], 'executions', '图 8 执行记录', '最多保留最近 100 次模拟记录。模拟完成不代表真实环境中的执行成功。')
page('9 阅读和编辑运维知识', '知识文档包含操作规程与技术指南，支持列表搜索、正文阅读、编辑和新建。', ['进入运维知识，搜索高 CPU 排查指南并打开。', '点击编辑文档，修改标题或正文后保存；刷新确认内容保留。', '编辑时选择取消，确认原文不变。', '返回列表，点击新建知识，填写必填标题和正文后保存。'], 'knowledge', '图 9 知识文档', '内置文档有中英文。编辑只修改开始编辑时的语言；用户新建内容不自动翻译。浏览器存储失败时请保留文本副本。')
page('10 演示高 CPU 处置闭环', '此场景演示订单服务资源、告警、知识、AI 和流程之间的关联。开始前先恢复示例，确保 INC-1201 为待处置。', ['从运维总览打开订单服务 CPU 持续高于阈值。', '确认告警，进入关联资源 order-api-01 查看 94% CPU 快照。', '返回告警详情，点击 AI 辅助分析，选择高 CPU 应急排查。', '查看关联知识，再通过前往关联流程模拟执行高 CPU 流程。', '核对执行记录后回到 INC-1201，确认关闭告警。'], 'assistant', '图 10 高 CPU 场景分析', '告警变为已关闭，执行记录增加一条。系统未执行真实限流、扩容或进程采样，资源 CPU 快照不会随告警关闭而降低。')
page('11 演示接口延迟分析', '此场景展示支付接口异常与连接池、慢查询分析建议的关联。开始前恢复示例数据。', ['进入告警中心，搜索 INC-1202 并打开。', '确认关联支付资源，点击确认告警。', '点击 AI 辅助分析，选择接口延迟诊断，阅读模拟建议。', '打开关联知识，随后进入接口延迟诊断流程并确认模拟执行。', '在执行记录核对流程与时间，再按演示需要关闭告警。'], 'automation', '图 11 接口延迟诊断流程', '生成接口延迟诊断的模拟执行记录。本版不采集真实链路、数据库连接或慢查询。')
page('12 演示每日健康巡检', '此场景展示从企业总览到标准巡检记录的演示路径。', ['从总览查看资源和待跟进风险数量。', '进入 AI 运维助手，选择每日健康巡检。', '阅读覆盖资源、需要关注资源与未关闭告警的模拟汇总。', '进入关联流程，确认模拟执行每日健康巡检。', '打开执行记录，核对新增记录的四个步骤。'], 'overview', '图 12 巡检前的运行态势', '形成一条巡检模拟记录，保留追溯入口。本版仅记录模拟流程完成，不生成可下载的真实巡检报告。')
page('13 个性化设置和恢复示例', '右上角账户菜单集中放置语言、主题、恢复示例、平台说明和退出。', ['打开账户菜单，选择切换 English，观察标题、按钮和状态文案。', '再次打开菜单，切换中文；选择切换深浅主题查看外观。', '恢复示例数据前先保存需要保留的知识文本。', '点击恢复示例数据，选择取消不会修改状态；确认后替换告警、知识和记录并重新生成时间快照。'], 'account', '图 13 账户菜单和恢复入口', '默认浅色，主题和语言在当前浏览器持久保存。恢复示例只重置演示数据，不重置语言或主题。')
page('14 常见问题与平台说明', '平台说明展示运行版本、数据边界、开源声明以及手册和对应源码下载。', ['页面显示存储警告时，先复制知识正文到安全位置，再尝试恢复示例或检查浏览器站点存储。', '换浏览器或访问地址后看不到旧数据属于预期行为；本版没有多人共享数据服务。', '刷新页面后需要重新进入演示时，点击进入演示即可；站点数据未清除时已保存内容仍可读取。', '遇到记录不存在，可能是此前恢复示例删除了自建记录；返回总览重新开始。', '平台说明保留许可证、第三方声明、上游归属和对应源码入口；下载操作手册获得本文件。'], 'about', '图 14 平台说明与下载入口', '本平台为独立定制演示版本，基于 Ongrid 开源项目，使用 AGPL-3.0-only，不代表上游官方产品或背书。')
args.output.parent.mkdir(parents=True, exist_ok=True)
doc.save(args.output)
print(args.output)
